"""In-memory parsing: user filenames never become filesystem paths."""

from io import BytesIO
from pathlib import Path
from src.models import Document, stable_id

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TOTAL_BYTES = 30 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


def load_bytes(name: str, data: bytes) -> list[Document]:
    name = name.replace("\\", "/").rsplit("/", 1)[-1]
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported document type: {suffix}")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Each document must be at most 10 MiB.")
    if suffix in {".txt", ".md"}:
        pages = [data.decode("utf-8-sig")]
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported.")
        pages = [page.extract_text() or "" for page in reader.pages]
    else:
        from docx import Document as WordDocument

        doc = WordDocument(BytesIO(data))
        pages = [
            "\n".join(
                [p.text for p in doc.paragraphs]
                + [
                    " | ".join(cell.text for cell in row.cells)
                    for table in doc.tables
                    for row in table.rows
                ]
            )
        ]
    if sum(map(len, pages)) > MAX_TEXT_CHARS:
        raise ValueError("Extracted text exceeds the 2 million character limit.")
    documents = [
        Document(stable_id(name, str(i), text), text.strip(), {"source": name, "page": i + 1})
        for i, text in enumerate(pages)
        if text.strip()
    ]
    if not documents:
        raise ValueError("No text found. Scanned PDFs need OCR before upload.")
    return documents


def load_documents(paths) -> list[Document]:
    paths = [Path(p) for p in paths]
    if sum(p.stat().st_size for p in paths) > MAX_TOTAL_BYTES:
        raise ValueError("Total uploads must be at most 30 MiB.")
    return [doc for path in paths for doc in load_bytes(path.name, path.read_bytes())]
