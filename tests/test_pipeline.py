from io import BytesIO
from types import SimpleNamespace
import pytest
from src.generation.generator import AnswerSchema, OpenAIGenerator
from src.ingestion.chunker import chunk_documents
from src.ingestion.loaders import load_bytes
from src.models import Document, Hit
from src.pipeline import Config, RAGPipeline
from src.retrieval.hybrid import SearchIndex, reciprocal_rank_fusion
from src.evaluation.datasets import load_fixture
from src.evaluation.judge import AnswerJudge


def test_chunk_identity_preserves_provenance():
    a = Document("a", "one two three four five", {"source": "a.txt"})
    b = Document("b", a.text, {"source": "b.txt"})
    chunks = chunk_documents([a, b], 3, 1)
    assert len({c.id for c in chunks}) == len(chunks)
    assert chunks == chunk_documents([a, b], 3, 1)
    assert chunks[1].text == "three four five"
    assert chunks[0].metadata["parent_id"] == "a"


@pytest.mark.parametrize("size,overlap", [(0, 0), (3, 3), (3, -1)])
def test_invalid_chunk_config(size, overlap):
    with pytest.raises(ValueError):
        chunk_documents([], size, overlap)


def test_upload_name_and_unicode():
    docs = load_bytes("../../private.txt", "cafÃ© document".encode())
    assert docs[0].metadata["source"] == "private.txt"
    assert docs[0].text == "cafÃ© document"
    with pytest.raises(ValueError):
        load_bytes("empty.txt", b" ")
    with pytest.raises(ValueError):
        load_bytes("bad.exe", b"hello")


def test_docx_paragraphs_and_tables():
    from docx import Document as WordDocument

    doc = WordDocument()
    doc.add_paragraph("Handbook")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Retention: 30 days"
    stream = BytesIO()
    doc.save(stream)
    result = load_bytes("handbook.docx", stream.getvalue())
    assert "Retention: 30 days" in result[0].text


def test_fusion_does_not_merge_identical_text_from_different_sources():
    a, b = Document("a", "same"), Document("b", "same")
    result = reciprocal_rank_fusion([[Hit(a, 1), Hit(b, 1)], [Hit(b, 1)]])
    assert [h.document.id for h in result] == ["b", "a"]


def test_retrieval_no_match_and_invalid_modes():
    index = SearchIndex(
        [Document("a", "apples fruit"), Document("b", "steel metal"), Document("c", "ocean water")]
    )
    assert index.search("apples", "bm25")[0].document.id == "a"
    assert index.search("quasar", "tfidf") == []
    with pytest.raises(ValueError):
        index.search("fruit", "dense")
    with pytest.raises(ValueError):
        index.search("fruit", "typo")


def test_instances_are_isolated_and_index_replacement_is_atomic():
    a, b = RAGPipeline(), RAGPipeline()
    docs, _, _ = load_fixture()
    a.index_records(docs, chunk=False)
    assert b.index is None
    with pytest.raises(ValueError):
        a.index_records([])
    assert a.query("API keys rotated")["citations"] == ["keys"]
    assert b.query("API keys rotated")["abstained"]
    a.index_records([Document("unique", "only a new document")], chunk=False)
    assert len(a.index.documents) == 1


@pytest.mark.parametrize("question", ["", " ", "a" * 4001])
def test_input_bounds(question):
    with pytest.raises(ValueError):
        RAGPipeline().query(question)


def test_config_validation():
    with pytest.raises(ValueError):
        Config(top_k=20, candidate_k=5)
    with pytest.raises(ValueError):
        Config(generation="unknown")


class FakeClient:
    def __init__(self, parsed):
        self.parsed = parsed
        self.responses = self
        self.arguments = None

    def parse(self, **kwargs):
        self.arguments = kwargs
        return SimpleNamespace(
            output_parsed=self.parsed, usage=SimpleNamespace(input_tokens=12, output_tokens=3)
        )


def test_structured_generation_validation_and_usage():
    fake = FakeClient(AnswerSchema(answer="30 days", citations=["a"], abstained=False))
    generator = OpenAIGenerator(client=fake)
    result = generator.generate("Retention?", [Hit(Document("a", "30 days"), 1)])
    assert result.usage == {"input_tokens": 12, "output_tokens": 3}
    assert fake.arguments["store"] is False
    fake.parsed = AnswerSchema(answer="unsupported", citations=["fake"], abstained=False)
    with pytest.raises(ValueError, match="citation"):
        generator.generate("Retention?", [Hit(Document("a", "30 days"), 1)])


def test_model_refusal_is_not_a_success():
    generator = OpenAIGenerator(client=FakeClient(None))
    with pytest.raises(ValueError, match="refusal"):
        generator.generate("Retention?", [Hit(Document("a", "30 days"), 1)])


def test_judge_missing_output_fails():
    judge = AnswerJudge(client=FakeClient(None))
    with pytest.raises(ValueError):
        judge.evaluate("q", {"answer": "a", "citations": [], "sources": []}, "reference")


def test_single_document_upload_is_searchable():
    index = SearchIndex([Document("only", "API keys rotate every 90 days")])
    assert index.search("API keys", "bm25")[0].document.id == "only"
    assert index.search("quasar", "bm25") == []


def test_extractive_abstains_on_stopword_only_overlap():
    from src.generation.generator import ExtractiveGenerator

    result = ExtractiveGenerator().generate(
        "Who won the chess championship?",
        [Hit(Document("a", "The database backup runs monthly."), 1)],
    )
    assert result.abstained
